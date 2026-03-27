# core/report_exporter.py
# ============================================================
#  LogiCheck — Exportador de Reportes (Excel y PDF básico)
# ============================================================

import os
import datetime


# ─────────────────────────────────────────────────────────────
#  EXCEL (openpyxl — ya instalado en el venv)
# ─────────────────────────────────────────────────────────────
def export_excel(audit_data: dict, output_path: str) -> bool:
    """
    Genera un reporte .xlsx con los datos de la auditoría.
    Retorna True si tuvo éxito.

    audit_data keys:
        factura_no, cliente, video_nombre, fecha,
        conteo_ia    {mat: n}, conteo_factura {mat: n},
        discrepancias {mat: diff}, vehiculo, resultado, usuario
    """
    try:
        import openpyxl
        from openpyxl.styles import (
            Font, PatternFill, Alignment, Border, Side
        )
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Reporte de Auditoría"

        # ── Paleta de colores ──
        azul_osc   = "FF1E1E2E"
        azul_med   = "FF313244"
        azul_claro = "FF89B4FA"
        verde      = "FFA6E3A1"
        rojo       = "FFF38BA8"
        amarillo   = "FFF9E2AF"
        blanco     = "FFCDD6F4"
        gris       = "FF6C7086"

        def apply_header(cell, text, bg=azul_osc, fg=azul_claro, bold=True, size=11):
            cell.value = text
            cell.font  = Font(color=fg, bold=bold, size=size, name="Segoe UI")
            cell.fill  = PatternFill("solid", fgColor=bg)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        def apply_cell(cell, text, bg=azul_med, fg=blanco, bold=False, align="left"):
            cell.value = text
            cell.font  = Font(color=fg, bold=bold, size=10, name="Segoe UI")
            cell.fill  = PatternFill("solid", fgColor=bg)
            cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)

        thin = Border(
            left  =Side(style="thin", color=azul_med),
            right =Side(style="thin", color=azul_med),
            top   =Side(style="thin", color=azul_med),
            bottom=Side(style="thin", color=azul_med),
        )

        def border_row(row, cols):
            for col in range(1, cols + 1):
                ws.cell(row=row, column=col).border = thin

        # ── ENCABEZADO PRINCIPAL ──
        ws.merge_cells("A1:F1")
        apply_header(ws["A1"], "LogiCheck — Reporte de Auditoría Logística",
                     bg=azul_osc, fg=azul_claro, size=14)
        ws.row_dimensions[1].height = 32

        ws.merge_cells("A2:F2")
        apply_cell(ws["A2"], f"Generado el {datetime.datetime.now().strftime('%d/%m/%Y %H:%M')}",
                   bg=azul_osc, fg=gris, align="center")
        ws.row_dimensions[2].height = 18

        # ── METADATA ──
        meta = [
            ("Factura Nº",  audit_data.get("factura_no",     "—")),
            ("Cliente",     audit_data.get("cliente",         "—")),
            ("Video",       audit_data.get("video_nombre",    "—")),
            ("Fecha",       audit_data.get("fecha",           "—")),
            ("Operador",    audit_data.get("usuario",         "—")),
            ("Vehículo",    audit_data.get("vehiculo",        "—")),
            ("Resultado",   audit_data.get("resultado",       "—")),
        ]

        ws.append([])  # fila vacía
        r = 4
        apply_header(ws.cell(r, 1), "Campo",  bg=azul_osc)
        apply_header(ws.cell(r, 2), "Valor",  bg=azul_osc)
        ws.merge_cells(f"B{r}:F{r}")
        border_row(r, 6)
        ws.row_dimensions[r].height = 22
        r += 1

        for key, val in meta:
            res_val = str(val)
            apply_cell(ws.cell(r, 1), key, bold=True)
            c = ws.cell(r, 2)
            # Resaltar resultado
            if key == "Resultado":
                if "CONFORME" in res_val:
                    apply_cell(c, res_val, fg=verde, bold=True, align="center")
                elif "DISCREPANCIA" in res_val:
                    apply_cell(c, res_val, fg=rojo, bold=True, align="center")
                else:
                    apply_cell(c, res_val)
            else:
                apply_cell(c, res_val)
            ws.merge_cells(f"B{r}:F{r}")
            border_row(r, 6)
            ws.row_dimensions[r].height = 20
            r += 1

        # ── TABLA DE CONTEO ──
        r += 1
        headers = ["Material", "Conteo Factura", "Conteo IA", "Diferencia", "Estado"]
        for ci, h in enumerate(headers, 1):
            cell = ws.cell(r, ci)
            apply_header(cell, h, bg=azul_osc)
            border_row(r, 5)
        ws.row_dimensions[r].height = 22
        r += 1

        conteo_ia      = audit_data.get("conteo_ia", {})
        conteo_factura = audit_data.get("conteo_factura", {})
        mats = sorted(set(list(conteo_ia.keys()) + list(conteo_factura.keys())))

        for mat in mats:
            ia_v  = int(conteo_ia.get(mat, 0))
            fac_v = int(conteo_factura.get(mat, 0))
            diff  = ia_v - fac_v
            if diff == 0:
                estado_txt = "✓ Conforme"
                estado_col = verde
            elif diff > 0:
                estado_txt = f"⚠ +{diff} exceso"
                estado_col = amarillo
            else:
                estado_txt = f"⚠ {diff} faltante"
                estado_col = rojo

            apply_cell(ws.cell(r, 1), mat, bold=True)
            apply_cell(ws.cell(r, 2), str(fac_v), align="center")
            apply_cell(ws.cell(r, 3), str(ia_v),  align="center")
            apply_cell(ws.cell(r, 4), str(diff),   align="center",
                       fg=(verde if diff == 0 else rojo))
            apply_cell(ws.cell(r, 5), estado_txt, fg=estado_col, bold=True, align="center")
            border_row(r, 5)
            ws.row_dimensions[r].height = 20
            r += 1

        # ── Anchos de columnas ──
        col_widths = [22, 18, 14, 12, 20, 20]
        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

        # ── EVIDENCIA VISUAL ──
        capturas = audit_data.get("capturas", [])
        if capturas:
            r += 1
            apply_header(ws.cell(r, 1), "Evidencia Visual (Capturas)", bg=azul_osc)
            ws.merge_cells(f"A{r}:F{r}")
            r += 1
            
            from openpyxl.drawing.image import Image as XLImage
            for img_path in capturas:
                if os.path.exists(img_path):
                    try:
                        img = XLImage(img_path)
                        # Redimensionar conservando proporción (aprox 400px ancho)
                        ratio = 400.0 / float(max(1, img.width))
                        img.width = 400
                        img.height = int(img.height * ratio)
                        ws.add_image(img, f"B{r}")
                        # Avanzar filas basándonos en la altura (aprox 1 fila = 15px)
                        r += int(img.height / 15) + 2
                    except Exception as e:
                        print(f"Error insertando imagen en Excel: {e}")

        ws.sheet_view.showGridLines = False
        wb.save(output_path)
        return True

    except Exception as e:
        print(f"[EXPORTER] Error exportando Excel: {e}")
        return False


# ─────────────────────────────────────────────────────────────
#  PDF (usando fpdf2)
# ─────────────────────────────────────────────────────────────
def export_pdf(audit_data: dict, output_path: str) -> bool:
    """
    Genera un reporte .pdf profesional usando fpdf2.
    """
    try:
        from fpdf import FPDF

        class ReportPDF(FPDF):
            def header(self):
                # Fondo oscuro para el encabezado
                self.set_fill_color(30, 30, 46) 
                self.rect(0, 0, 210, 40, "F")
                
                self.set_y(10)
                self.set_font("Helvetica", "B", 22)
                self.set_text_color(137, 180, 250) # Azul LogiCheck
                self.cell(0, 10, "LOGICHECK", ln=True, align="C")
                
                self.set_font("Helvetica", "", 10)
                self.set_text_color(166, 227, 161) # Verde
                self.cell(0, 8, "AUDITORÍA DE DESPACHO E INTELIGENCIA ARTIFICIAL", ln=True, align="C")
                self.ln(12)

            def footer(self):
                self.set_y(-20)
                self.set_font("Helvetica", "I", 8)
                self.set_text_color(108, 112, 134)
                self.cell(0, 10, f"Documento generado automáticamente por LogiCheck v1.2 - Página {self.page_no()}", align="C")

        pdf = ReportPDF()
        pdf.add_page()
        pdf.set_auto_page_break(True, margin=20)

        # Colores
        MED   = (49, 50, 68)
        BLUE  = (137, 180, 250)
        GREEN = (166, 227, 161)
        RED   = (243, 139, 168)
        YELL  = (249, 226, 175)
        WHITE = (205, 214, 244)
        GRAY  = (108, 112, 134)
        BG    = (30, 30, 46)

        # Fondo general de la página (opcional, pero se ve premium)
        pdf.set_fill_color(*BG)
        # pdf.rect(0, 40, 210, 257, "F") # Desactivado por ahora para legibilidad al imprimir

        def section_title(text):
            pdf.ln(5)
            pdf.set_fill_color(*MED)
            pdf.set_text_color(*BLUE)
            pdf.set_font("Helvetica", "B", 12)
            pdf.cell(0, 10, f"  {text}", ln=True, fill=True)
            pdf.ln(3)

        def info_row(key, val, val_color=None):
            pdf.set_text_color(*GRAY)
            pdf.set_font("Helvetica", "B", 10)
            pdf.cell(50, 8, f" {key}:")
            pdf.set_text_color(*(val_color or (30, 30, 30)))
            pdf.set_font("Helvetica", "", 10)
            pdf.cell(0, 8, str(val), ln=True)

        # 1. Info General
        section_title("Resumen del Despacho")
        info_row("Factura No", audit_data.get("factura_no", "—"))
        info_row("Cliente",    audit_data.get("cliente",    "—"))
        info_row("Operador",   audit_data.get("usuario",    "—"))
        info_row("Fecha",      audit_data.get("fecha",      "—"))
        info_row("Vehículo",   audit_data.get("vehiculo",   "No asignado"))
        
        res = audit_data.get("resultado", "")
        info_row("Resultado", res, val_color=(GREEN if "CONFORME" in res else RED))

        # 2. Tabla de Comparación
        section_title("Análisis Detallado IA vs Factura")
        
        # Headers de tabla
        pdf.set_fill_color(*MED)
        pdf.set_text_color(*WHITE)
        pdf.set_font("Helvetica", "B", 10)
        pdf.cell(70, 10, " Material", fill=True)
        pdf.cell(30, 10, "Factura", fill=True, align="C")
        pdf.cell(30, 10, "IA", fill=True, align="C")
        pdf.cell(30, 10, "Dif.", fill=True, align="C")
        pdf.cell(30, 10, "Estado", fill=True, align="C")
        pdf.ln()

        conteo_ia      = audit_data.get("conteo_ia", {})
        conteo_factura = audit_data.get("conteo_factura", {})
        mats = sorted(set(list(conteo_ia.keys()) + list(conteo_factura.keys())))

        pdf.set_text_color(30, 30, 30)
        pdf.set_font("Helvetica", "", 10)
        
        for i, mat in enumerate(mats):
            ia_v  = int(conteo_ia.get(mat, 0))
            fac_v = int(conteo_factura.get(mat, 0))
            diff  = ia_v - fac_v
            
            # Alternar color de fondo de fila
            if i % 2 == 0:
                pdf.set_fill_color(245, 245, 250)
            else:
                pdf.set_fill_color(255, 255, 255)
            
            pdf.cell(70, 9, f" {mat}", fill=True)
            pdf.cell(30, 9, str(fac_v), fill=True, align="C")
            pdf.cell(30, 9, str(ia_v), fill=True, align="C")
            
            diff_col = GREEN if diff == 0 else (RED if diff < 0 else YELL)
            pdf.set_text_color(*diff_col)
            pdf.cell(30, 9, f"{diff:+d}", fill=True, align="C")
            
            pdf.set_font("Helvetica", "B", 9)
            estado = "OK" if diff == 0 else "ERROR"
            pdf.cell(30, 9, estado, fill=True, align="C")
            
            pdf.set_text_color(30, 30, 30)
            pdf.set_font("Helvetica", "", 10)
            pdf.ln()

        # 3. Evidencias Visuales (Capturas)
        capturas = audit_data.get("capturas", [])
        if capturas:
            section_title("Evidencia Visual (Capturas Automáticas)")
            for img_path in capturas:
                if os.path.exists(img_path):
                    # Nueva página si no hay espacio (aprox 120mm)
                    if pdf.get_y() > 160:
                        pdf.add_page()
                    try:
                        pdf.image(img_path, w=150, x=(210-150)/2)
                        pdf.ln(5)
                    except Exception as e:
                        print(f"Error insertando imagen en PDF: {e}")

        pdf.output(output_path)
        return True

    except Exception as e:
        print(f"[EXPORTER] Error exportando PDF: {e}")
        return False


def _export_txt_fallback(audit_data: dict, output_path: str) -> bool:
    """Exporta un resumen de texto si fpdf2 no está disponible."""
    try:
        lines = [
            "=" * 60,
            "  LogiCheck — Reporte de Auditoría Logística",
            f"  Generado: {datetime.datetime.now().strftime('%d/%m/%Y %H:%M')}",
            "=" * 60,
            "",
            f"  Factura:   {audit_data.get('factura_no', '—')}",
            f"  Cliente:   {audit_data.get('cliente', '—')}",
            f"  Video:     {audit_data.get('video_nombre', '—')}",
            f"  Fecha:     {audit_data.get('fecha', '—')}",
            f"  Operador:  {audit_data.get('usuario', '—')}",
            f"  Vehículo:  {audit_data.get('vehiculo', '—')}",
            f"  Resultado: {audit_data.get('resultado', '—')}",
            "",
            "-" * 60,
            f"  {'Material':<25} {'Factura':>8} {'IA':>6} {'Dif':>6}",
            "-" * 60,
        ]
        conteo_ia      = audit_data.get("conteo_ia", {})
        conteo_factura = audit_data.get("conteo_factura", {})
        mats = sorted(set(list(conteo_ia.keys()) + list(conteo_factura.keys())))
        for mat in mats:
            ia_v  = int(conteo_ia.get(mat, 0))
            fac_v = int(conteo_factura.get(mat, 0))
            diff  = ia_v - fac_v
            lines.append(f"  {mat:<25} {fac_v:>8} {ia_v:>6} {diff:>+6}")
        lines += ["=" * 60, "  (Nota: fpdf2 no instalado — exportado como texto)", "=" * 60]

        txt_path = output_path.replace(".pdf", ".txt")
        with open(txt_path, "w", encoding="utf-8") as f:
            f.write("\n".join(lines))
        return True
    except Exception as e:
        print(f"[EXPORTER] Error en fallback txt: {e}")
        return False
