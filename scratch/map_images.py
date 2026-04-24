"""Reads existing docx and maps each image to its paragraph context."""
from docx import Document
from lxml import etree

doc = Document(r'C:\Users\samuv\Desktop\Programas_mios\LogiCheck\TESIS\Manual de Usuario\Manual_Usuario_LogiCheck_v1.docx')

# Build a mapping: blip embed id -> image filename
rels_map = {}
img_counter = 0
for rel_id, rel in doc.part.rels.items():
    if 'image' in rel.reltype:
        img_counter += 1
        import os
        ext = os.path.splitext(rel.target_ref)[1]
        rels_map[rel_id] = f'img_{img_counter:02d}{ext}'

# Now walk paragraphs and find images
ns = {
    'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
    'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing',
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    'pic': 'http://schemas.openxmlformats.org/drawingml/2006/picture',
}

print("=== IMAGE-TO-PARAGRAPH MAPPING ===\n")
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    style = p.style.name if p.style else 'Normal'
    
    # Find blip elements (embedded images)
    blips = p._element.findall('.//a:blip', ns)
    if blips:
        for blip in blips:
            embed_id = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
            img_file = rels_map.get(embed_id, '???')
            # Get prev paragraph text for context
            prev_text = doc.paragraphs[i-1].text.strip() if i > 0 else ''
            next_text = doc.paragraphs[i+1].text.strip() if i < len(doc.paragraphs)-1 else ''
            print(f"Para [{i}] ({style}): IMAGE={img_file}")
            print(f"  Text: '{text[:100]}'")
            print(f"  Prev: '{prev_text[:100]}'")
            print(f"  Next: '{next_text[:100]}'")
            print()
